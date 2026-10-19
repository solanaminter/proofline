#!/usr/bin/env python3
"""Build the Proofline pitch deck (PPTX) — 7 slides, 16:9, IBM-blue accent.
Editable source for docs/proofline-pitch-deck.pdf. Every claim verified against
README.md, lib/taskpack.ts, scripts/eval.ts, docs/DEMO_VIDEO_SCRIPT.md.
Run: python3 build_deck.py  -> writes proofline-pitch-deck.pptx in this dir.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Theme
BG = RGBColor(0x0B, 0x15, 0x26)       # dark navy
CARD = RGBColor(0x13, 0x22, 0x3A)     # card navy
BLUE = RGBColor(0x0F, 0x62, 0xFE)     # IBM blue accent
LTBLUE = RGBColor(0x78, 0xA9, 0xFF)   # light blue
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0xA6, 0xAD, 0xBA)     # cool gray 50
DGRAY = RGBColor(0x69, 0x70, 0x74)    # cool gray 60
AMBER = RGBColor(0xF1, 0xC2, 0x1C)    # pending/amber
FONT = "DejaVu Sans"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = BG
    return s


def tb(s, x, y, w, h):
    shp = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shp.text_frame.word_wrap = True
    return shp


def para(tf, text, size=20, bold=False, color=WHITE, align=PP_ALIGN.LEFT,
         font=FONT, italic=False, space_after=6):
    p = tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = font
    return p


def first(tf, text, size=20, bold=False, color=WHITE, align=PP_ALIGN.LEFT,
          font=FONT, italic=False, space_after=6):
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(space_after)
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = font
    return p


def header(s, kicker, title):
    t = tb(s, 0.7, 0.35, 11.9, 1.6)
    first(t.text_frame, kicker, size=15, bold=True, color=LTBLUE)
    para(t.text_frame, title, size=36, bold=True, color=WHITE)
    # accent rule
    shape = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.95),
                               Inches(0.9), Inches(0.07))
    shape.fill.solid(); shape.fill.fore_color.rgb = BLUE
    shape.line.fill.background()


def footer(s, rubric):
    t = tb(s, 0.7, 6.95, 11.9, 0.4)
    first(t.text_frame, rubric, size=11, color=DGRAY)


def bullets(s, items, x=0.7, y=2.35, w=11.9, size=19, gap=10, color=WHITE,
            lead_color=LTBLUE):
    t = tb(s, x, y, w, 4.3)
    tf = t.text_frame
    tf.word_wrap = True
    first_done = False
    for lead, rest in items:
        if not first_done:
            p = tf.paragraphs[0]
            first_done = True
        else:
            p = tf.add_paragraph()
        p.space_after = Pt(gap)
        p.level = 0
        r = p.add_run()
        r.text = ("▸ " + lead + "  ") if lead else "▸ "
        r.font.size = Pt(size); r.font.bold = bool(lead)
        r.font.color.rgb = lead_color; r.font.name = FONT
        r2 = p.add_run()
        r2.text = rest
        r2.font.size = Pt(size); r2.font.color.rgb = color; r2.font.name = FONT
    return t


def box(s, x, y, w, h, fill=CARD, border=None):
    shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                               Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    if border:
        shape.line.color.rgb = border; shape.line.width = Pt(2)
    else:
        shape.line.fill.background()
    return shape


def boxtext(s, x, y, w, h, title, body, tsize=17, bsize=15, accent=None):
    shape = box(s, x, y, w, h)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25); tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.18)
    first(tf, title, size=tsize, bold=True, color=accent or WHITE)
    para(tf, body, size=bsize, color=GRAY)
    return shape


RUB = {
    1: "Presentation & Clarity  ·  10%",
    2: "Tool Usage & Depth of Adoption  ·  30%",
    3: "Tool Usage & Depth of Adoption  ·  30%",
    4: "Innovation & Use Case Relevance  ·  30%",
    5: "Innovation & Use Case Relevance  ·  30%",
    6: "Functionality & Technical Execution  ·  30%",
    7: "Presentation & Clarity  ·  10%",
}

# ---------- Slide 1: Title ----------
s = slide()
t = tb(s, 0.7, 1.6, 11.9, 3.5)
first(t.text_frame, "PROOFLINE", size=72, bold=True, color=WHITE)
para(t.text_frame, "Evidence-first change planning for unfamiliar repositories",
     size=26, color=LTBLUE)
para(t.text_frame,
     "Paste a repo URL + a change request → impact graph, evidence-linked plan, "
     "and a Bob-ready task pack — in under 45 seconds.", size=17, color=GRAY)
t2 = tb(s, 0.7, 5.6, 11.9, 1.2)
first(t2.text_frame, "AngelHack  ·  Building with IBM Bob  ·  Theme 1: "
      "Explore, Fix, Build an unfamiliar app", size=16, bold=True, color=WHITE)
para(t2.text_frame, "Solana Minter  ·  solo  ·  Purple Castle Ventures",
     size=15, color=GRAY)
footer(s, RUB[1])

# ---------- Slide 2: Bob workflow (1/2) — the loop ----------
s = slide()
header(s, "IBM BOB WORKFLOW  ·  1/2", "Give Bob a map before the keys")
bullets(s, [
    ("The problem. ",
     "Coding agents edit code they've never seen and guess at the blast radius — "
     "confidently. Every wrong file touched is a regression waiting to happen."),
    ("The loop. ",
     "Proofline turns a change request into an auditable plan that IBM Bob "
     "executes file-by-file: repo URL + request → impact graph → evidence-linked "
     "plan → task pack → Bob implements with the guardrail "
     "\"do not edit files outside the plan\"."),
    ("Depth of adoption. ",
     "Bob is a core part of the workflow, not bolted on: every task pack ships a "
     "BOB_PROMPT.md written for Bob's Plan/Agent mode, with the evidence, test "
     "commands, and edit boundaries baked in."),
])
footer(s, RUB[2])

# ---------- Slide 3: Bob workflow (2/2) — honest evidence status ----------
s = slide()
header(s, "IBM BOB WORKFLOW  ·  2/2", "Evidence status — stated honestly")
bullets(s, [
    ("Verified in code. ",
     "Task-pack generation is unit-tested: every pack contains TASK.md, "
     "CONTEXT.json, TEST_PLAN.md, and BOB_PROMPT.md with the file-by-file "
     "instructions and the \"do not edit files outside the plan\" guardrail."),
    ("What the pack looks like. ",
     "BOB_PROMPT.md = scope (files in/out), the evidence behind each edit, the "
     "test commands to run, and explicit stop conditions — everything Bob needs "
     "to execute without improvising."),
])
# pending box
shape = box(s, 0.7, 5.0, 11.93, 1.55, fill=RGBColor(0x2B, 0x21, 0x08),
            border=AMBER)
tf = shape.text_frame
tf.word_wrap = True
tf.margin_left = Inches(0.3); tf.margin_right = Inches(0.3)
tf.margin_top = Inches(0.2)
first(tf, "⏳  PENDING — Bob session evidence",
      size=18, bold=True, color=AMBER)
para(tf, "Live session exports (bob-sessions/) are finalized once AngelHack "
         "registration unlocks IBM Bob IDE access. The task-pack → Bob loop is "
         "designed and code-verified today; nothing here is simulated or "
         "fabricated.", size=15, color=GRAY)
footer(s, RUB[3])

# ---------- Slide 4: Use case — canonical demo ----------
s = slide()
header(s, "USE CASE", "The demo: a real change, end to end")
bullets(s, [
    ("Setup. ",
     "Built-in demo repository (a small TypeScript API client) loads instantly — "
     "no network needed on camera."),
    ("Request. ",
     "Change brief: \"Handle API errors with retries\" — one of the verified "
     "eval briefs, so the plan on screen is the plan the eval suite scores."),
    ("In under 45 seconds. ",
     "Proofline parses the codebase and shows: an impact graph of every file the "
     "change could touch; an evidence-linked change plan where each file carries "
     "its reason; a risk list; a test strategy; and the downloadable task pack."),
])
footer(s, RUB[4])

# ---------- Slide 5: Innovation — deterministic evidence ranking ----------
s = slide()
header(s, "INNOVATION", "Deterministic ranking, auditable by construction")
bullets(s, [
    ("No LLM, no randomness. ",
     "ts-morph parses the repo into an AST of imports, symbols, tests, and "
     "config. Every file is scored by evidence weight — run it twice, get the "
     "same plan."),
    ("Evidence weights. ",
     "Lexical match 55  ·  import distance 20  ·  symbol references 12  ·  "
     "tests 8  ·  configuration 5."),
    ("Honest uncertainty. ",
     "Weak signals are labeled \"candidate\" — never dressed up as certainty. "
     "Every recommendation traces to an import, symbol, test, or config "
     "reference. Analysis lives in request memory only; no repos are stored."),
])
footer(s, RUB[5])

# ---------- Slide 6: Technical execution ----------
s = slide()
header(s, "TECHNICAL EXECUTION", "Architecture, stack, and proof")
# pipeline boxes
steps = [("repo URL", "public zipball\n200 files / 2 MB cap"),
         ("ts-morph\nparse", "imports · symbols\ntests · config"),
         ("evidence\nrank", "55 / 20 / 12\n8 / 5"),
         ("plan", "graph + risks\n+ test strategy"),
         ("task pack", "TASK.md · CONTEXT.json\nTEST_PLAN.md · BOB_PROMPT.md")]
x = 0.7
for i, (title, body) in enumerate(steps):
    boxtext(s, x, 2.45, 2.05, 1.7, title, body, tsize=16, bsize=13,
            accent=LTBLUE)
    if i < len(steps) - 1:
        at = tb(s, x + 2.05, 3.05, 0.44, 0.5)
        first(at.text_frame, "→", size=24, bold=True, color=BLUE,
              align=PP_ALIGN.CENTER)
    x += 2.49
bullets(s, [
    ("Stack. ",
     "Next.js 15 + TypeScript · ts-morph (server AST) · Cytoscape.js (client "
     "graph) · Tailwind CSS v4. Vercel-deployable."),
    ("Verification. ",
     "Evals 10/10 — five change briefs × two fixtures, top-5 accuracy + "
     "runtime. Unit tests 9/9, green."),
    ("Honest limits. ",
     "TypeScript/JavaScript public repos only; live analysis targets <45s with "
     "a fallback ladder (live → cached fixture → error, never a blank page). "
     "Ranking is heuristic evidence aggregation — low-signal files say so."),
], y=4.55)
footer(s, RUB[6])

# ---------- Slide 7: Roadmap + learnings ----------
s = slide()
header(s, "ROADMAP + LEARNINGS", "What's next, and what we learned")
t = tb(s, 0.7, 2.35, 5.7, 4.2)
first(t.text_frame, "Roadmap", size=22, bold=True, color=LTBLUE)
for item in ["Bob session exports under bob-sessions/ once IDE access lands",
             "Live deployment of the planner for judges to try",
             "Multi-language parsing beyond TypeScript/JavaScript",
             "Session evidence feeding back into ranking weights"]:
    para(t.text_frame, "▸  " + item, size=16, color=WHITE)
t2 = tb(s, 6.9, 2.35, 5.7, 4.2)
first(t2.text_frame, "Learnings", size=22, bold=True, color=LTBLUE)
for item in ["Deterministic beats LLM for auditability — the plan must be "
             "reproducible",
             "Guardrails make agents trustworthy: scope the edits, and the "
             "agent stops improvising",
             "Evidence labels beat confidence theater — \"candidate\" is a "
             "feature, not a failure"]:
    para(t2.text_frame, "▸  " + item, size=16, color=WHITE)
footer(s, RUB[7])

prs.save("proofline-pitch-deck.pptx")
print("saved proofline-pitch-deck.pptx:", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
